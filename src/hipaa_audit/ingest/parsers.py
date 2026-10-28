"""Turn an uploaded PDF, DOCX, TXT or Markdown file into ordered blocks with page and heading information.

Parsing is deliberately defensive: the file is untrusted input. File type is checked by extension and by magic bytes, size
and page counts are bounded, encrypted and scanned PDFs are rejected with a clear message, and DOCX archives are checked for
decompression bombs before they are opened.
"""

from __future__ import annotations

import io
import logging
import os
import re
import zipfile

from hipaa_audit.models import Block

log = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".markdown")
MIN_BODY_WORDS = 20
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_MEMBERS = 5000


class UnsupportedDocument(ValueError):
    """The upload cannot be turned into text (wrong type, too large, encrypted, scanned or empty)."""


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")
_NUMBERED = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+([A-Z][^\n]*)$")
_BULLET = re.compile(r"^\s*(?:[•▪●◦\-–*]|\(?\d{1,2}[.)]|\(?[a-z][.)])\s+")
_SENTENCE_WORDS = {"must", "shall", "will", "should", "may", "is", "are", "be", "can", "has", "have"}


class _Outline:
    """Tracks the heading path while a document is read in order."""

    def __init__(self) -> None:
        self.stack: list[tuple[int, str]] = []

    def push(self, level: int, title: str) -> None:
        while self.stack and self.stack[-1][0] >= level:
            self.stack.pop()
        self.stack.append((level, title))

    @property
    def path(self) -> str:
        return " > ".join(title for _, title in self.stack)


def heading_of(line: str, *, previous: str | None = None) -> tuple[int, str] | None:
    """Heuristic heading detection for plain text and PDF lines. Returns (level, title) or None.

    ``previous`` is the line before this one when it belongs to the same run of text (``None`` at the start of a page or after a
    heading or a blank line). A single capitalised word on a line after one that does not end a sentence is an acronym that wrapped
    or a table cell, not a heading, so it is left in the text.
    """
    s = line.strip()
    if not s or len(s) > 100:
        return None
    m = _MD_HEADING.match(s)
    if m:
        return len(m.group(1)), m.group(2).strip()
    if s.endswith((".", ",", ";", ":")):
        return None
    words = s.split()
    m = _NUMBERED.match(s)
    if m and len(words) <= 10 and not any(w.lower() in _SENTENCE_WORDS for w in words):
        return m.group(1).count(".") + 2, s  # level 1 is reserved for a document title
    letters = [ch for ch in s if ch.isalpha()]
    if len(letters) >= 4 and s == s.upper() and len(words) <= 10:
        if len(words) == 1 and previous is not None and not previous.strip().endswith((".", "!", "?", ":", ";")):
            return None
        return 1, s
    return None


def _paragraphs(text: str, page: int | None, outline: _Outline) -> list[Block]:
    """Group lines into paragraphs (blank lines, headings and bullets start a new one)."""
    blocks: list[Block] = []
    current: list[str] = []

    def flush() -> None:
        if current:
            body = " ".join(current).strip()
            if body:
                blocks.append(Block(text=body, page=page, section=outline.path))
            current.clear()

    previous: str | None = None  # the line before this one, while it is part of the same run of text
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
            previous = None
            continue
        found = heading_of(line, previous=previous)
        if found:
            flush()
            level, title = found
            outline.push(level, title)
            blocks.append(Block(text=title, page=page, section=outline.path, is_heading=True))
            previous = None
            continue
        if _BULLET.match(raw):
            flush()
        current.append(line)
        previous = line
    flush()
    return blocks


def _parse_text(data: bytes) -> list[Block]:
    if b"\x00" in data[:4096]:
        raise UnsupportedDocument("the file looks like a binary file, not text")
    text = data.decode("utf-8-sig", errors="replace")
    return _paragraphs(text, None, _Outline())


def _parse_pdf(data: bytes, max_pages: int) -> list[Block]:
    if b"%PDF-" not in data[:1024]:
        raise UnsupportedDocument("the file is not a PDF")
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedDocument("the PDF could not be read") from exc
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            raise UnsupportedDocument("the PDF is password protected")
    n_pages = len(reader.pages)
    if n_pages > max_pages:
        raise UnsupportedDocument(f"the PDF has {n_pages} pages; the limit is {max_pages}")
    outline = _Outline()
    blocks: list[Block] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        blocks.extend(_paragraphs(text, number, outline))
    return blocks


def _check_docx_archive(data: bytes) -> None:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise UnsupportedDocument("the file is not a DOCX document")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_DOCX_MEMBERS or sum(m.file_size for m in members) > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise UnsupportedDocument("the DOCX file is too large once unpacked")


def _parse_docx(data: bytes) -> list[Block]:
    _check_docx_archive(data)
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedDocument("the DOCX file could not be read") from exc
    outline = _Outline()
    blocks: list[Block] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"^Heading (\d)$", style)
            level = int(m.group(1)) if m else (0 if style == "Title" else None)
            if level is None:
                found = heading_of(text)
                level = found[0] if found and len(text.split()) <= 10 else None
            if level is not None:
                outline.push(max(level, 1), text)
                blocks.append(Block(text=text, section=outline.path, is_heading=True))
            else:
                blocks.append(Block(text=text, section=outline.path))
        elif tag == "tbl":
            table = Table(child, doc)
            for row in table.rows:
                seen: set[int] = set()
                cells: list[str] = []
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    value = " ".join(cell.text.split())
                    if value:
                        cells.append(value)
                if cells:
                    blocks.append(Block(text=" | ".join(cells), section=outline.path))
    return blocks


def parse_document(filename: str, data: bytes, *, max_pages: int = 200) -> list[Block]:
    """Parse an uploaded file. Raises ``UnsupportedDocument`` with a user-presentable message."""
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedDocument(f"unsupported file type '{ext or filename}'; use PDF, DOCX, TXT or Markdown")
    if not data:
        raise UnsupportedDocument("the file is empty")
    try:
        if ext == ".pdf":
            blocks = _parse_pdf(data, max_pages)
        elif ext == ".docx":
            blocks = _parse_docx(data)
        else:
            blocks = _parse_text(data)
    except UnsupportedDocument:
        raise
    except Exception as exc:  # a damaged file can break a parser in ways its authors did not anticipate
        log.warning("an upload with extension %s could not be parsed (%s)", ext, type(exc).__name__)
        what = {".pdf": "PDF", ".docx": "DOCX file"}.get(ext, "file")
        raise UnsupportedDocument(f"the {what} could not be read") from exc
    body_words = sum(len(b.text.split()) for b in blocks if not b.is_heading)
    if body_words < MIN_BODY_WORDS:
        raise UnsupportedDocument("no readable text found (a scanned PDF needs OCR first)")
    return blocks
