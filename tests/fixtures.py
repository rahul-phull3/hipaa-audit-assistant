"""Builders for small test documents. Nothing binary is committed to the repository."""

from __future__ import annotations

import io

from docx import Document

SAMPLE_POLICY = """ACCESS CONTROL POLICY

1. Purpose
This policy sets the rules for granting access to systems that hold electronic protected health information.

2. Password Management
- Passwords must be at least 14 characters long and must never be shared with anyone.
- Passwords are stored only in the approved password manager.

2.1 Automatic Logoff
Sessions are logged off automatically after 15 minutes of inactivity.
"""


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF. Each page is a list of text lines set in 11 pt Helvetica."""
    n_pages = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, lines in enumerate(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        ops = ["BT", "/F1 11 Tf", "14 TL", "50 740 Td"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({escaped}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Information Security Policy", level=1)
    doc.add_paragraph(
        "This policy applies to every workforce member who creates, receives, maintains or transmits electronic "
        "protected health information on behalf of the organization."
    )
    doc.add_heading("Access Control", level=2)
    doc.add_paragraph("Access to systems that hold ePHI is granted only after approval by the data owner.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Review"
    table.cell(0, 1).text = "Quarterly"
    table.cell(1, 0).text = "Owner"
    table.cell(1, 1).text = "Security official"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
